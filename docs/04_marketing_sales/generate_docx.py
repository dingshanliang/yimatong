# -*- coding: utf-8 -*-
"""
一码通功能清单 — 专业 Word 文档生成器
配色：深墨绿 #0D3B2E + 金 #C8952D + 浅底 #F5F1EA
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import copy

# ===== 品牌色 =====
INK = RGBColor(0x0D, 0x3B, 0x2E)       # 深墨绿
INK_HEX = "0D3B2E"
GOLD = RGBColor(0xC8, 0x95, 0x2D)      # 金
GOLD_HEX = "C8952D"
GOLD_SOFT_HEX = "EBD9B8"
PAPER_HEX = "F5F1EA"
RULE_HEX = "D4CDBF"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5C, 0x6B, 0x63)

# ===== 字体 =====
FONT_DISPLAY = "Microsoft YaHei"    # 标题
FONT_BODY = "Microsoft YaHei"       # 正文

doc = Document()

# ===== 页面设置 =====
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.5)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)

# ===== 默认样式 =====
style = doc.styles['Normal']
style.font.name = FONT_BODY
style.font.size = Pt(11)
style.font.color.rgb = DARK_TEXT
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.4

# 中文字体绑定
rpr = style.element.rPr
if rpr is None:
    rpr = parse_xml(f'<w:rPr {nsdecls("w")}></w:rPr>')
    style.element.append(rpr)
rFonts = rpr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
    rpr.insert(0, rFonts)
rFonts.set(qn('w:eastAsia'), FONT_BODY)


# ============================================================
# 辅助函数
# ============================================================

def set_cell_shading(cell, color_hex):
    """设置单元格背景色"""
    shading = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>'
    )
    cell._element.tcPr.append(shading)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """设置单元格内边距（twips）"""
    tcPr = cell._element.tcPr
    margins = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:start w:w="{left}" w:type="dxa"/>'
        f'  <w:end w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(margins)


def add_cell_text(cell, text, font_size=10.5, bold=False, color=DARK_TEXT, align=WD_ALIGN_PARAGRAPH.LEFT, font_name=None):
    """向单元格写入格式化文本"""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    fn = font_name or FONT_BODY
    run.font.name = fn
    r = run._element
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = parse_xml(f'<w:rPr {nsdecls("w")}></w:rPr>')
        r.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), fn)


def add_heading_1(text):
    """一级标题：深绿底色色块条 + 白字"""
    # 空行间距
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0

    # 色块标题
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, INK_HEX)
    set_cell_margins(cell, top=100, bottom=100, left=200, right=200)

    # 宽度撑满
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')

    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = FONT_DISPLAY
    rPr = run._element.find(qn('w:rPr'))
    rFonts = rPr.find(qn('w:rFonts'))
    rFonts.set(qn('w:eastAsia'), FONT_DISPLAY)

    # 色块宽度 = 页面宽度
    for row in table.rows:
        row.cells[0].width = Cm(15.7)

    return table


def add_heading_2(text):
    """二级标题：金色竖线 + 深绿文字"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)

    # 金色竖线
    run_bar = p.add_run("▎")
    run_bar.font.size = Pt(15)
    run_bar.font.color.rgb = GOLD
    run_bar.font.name = FONT_DISPLAY
    rPr = run_bar._element.find(qn('w:rPr'))
    if rPr is not None:
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is not None:
            rFonts.set(qn('w:eastAsia'), FONT_DISPLAY)

    run = p.add_run(f" {text}")
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = INK
    run.font.name = FONT_DISPLAY
    rPr = run._element.find(qn('w:rPr'))
    if rPr is not None:
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is not None:
            rFonts.set(qn('w:eastAsia'), FONT_DISPLAY)


def add_body_paragraph(text, indent_first=False, font_size=11):
    """正文段落"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    if indent_first:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.font.color.rgb = DARK_TEXT
    run.font.name = FONT_BODY
    rPr = run._element.find(qn('w:rPr'))
    rFonts = rPr.find(qn('w:rFonts'))
    rFonts.set(qn('w:eastAsia'), FONT_BODY)


def add_bullet(text, bold_prefix=None):
    """无序列表项"""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.left_indent = Cm(0.8)

    if bold_prefix:
        run_b = p.add_run(bold_prefix)
        run_b.font.size = Pt(11)
        run_b.font.bold = True
        run_b.font.color.rgb = INK
        run_b.font.name = FONT_BODY
        rPr = run_b._element.find(qn('w:rPr'))
        rFonts = rPr.find(qn('w:rFonts'))
        rFonts.set(qn('w:eastAsia'), FONT_BODY)

    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_TEXT
    run.font.name = FONT_BODY
    rPr = run._element.find(qn('w:rPr'))
    rFonts = rPr.find(qn('w:rFonts'))
    rFonts.set(qn('w:eastAsia'), FONT_BODY)


def add_note_block(text):
    """灰色注释框"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, PAPER_HEX)
    set_cell_margins(cell, top=120, bottom=120, left=200, right=200)

    # 左边金色边框
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none" w:sz="0" w:space="0"/>'
        f'  <w:bottom w:val="none" w:sz="0" w:space="0"/>'
        f'  <w:start w:val="single" w:sz="24" w:color="{GOLD_HEX}"/>'
        f'  <w:end w:val="none" w:sz="0" w:space="0"/>'
        f'</w:tcBorders>'
    )
    cell._element.tcPr.append(tcBorders)

    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.4
    run = p.add_run("说明：")
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = GOLD
    run.font.name = FONT_BODY
    rPr = run._element.find(qn('w:rPr'))
    rFonts = rPr.find(qn('w:rFonts'))
    rFonts.set(qn('w:eastAsia'), FONT_BODY)

    run2 = p.add_run(text)
    run2.font.size = Pt(10.5)
    run2.font.color.rgb = MUTED
    run2.font.name = FONT_BODY
    rPr2 = run2._element.find(qn('w:rPr'))
    rFonts2 = rPr2.find(qn('w:rFonts'))
    rFonts2.set(qn('w:eastAsia'), FONT_BODY)

    for row in table.rows:
        row.cells[0].width = Cm(15.7)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_feature_table(headers, rows):
    """
    通用功能表格生成器
    headers: [str, str, str]
    rows: [[功能, 业务结果, 层级], ...]
    """
    n_cols = len(headers)
    n_rows = len(rows) + 1

    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 表头
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        set_cell_shading(cell, INK_HEX)
        set_cell_margins(cell)
        add_cell_text(cell, h, font_size=10.5, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

    # 数据行
    for i, row_data in enumerate(rows):
        bg = "FFFFFF" if i % 2 == 0 else "F8F6F1"
        for j, val in enumerate(row_data):
            cell = table.cell(i + 1, j)
            set_cell_shading(cell, bg)
            set_cell_margins(cell)

            # 第三列（层级）特殊着色
            if j == 2:
                level_colors = {
                    "基础能力": INK,
                    "增强能力": GOLD,
                    "对接/定制": MUTED,
                    "项目服务": RGBColor(0x8B, 0x45, 0x13),
                }
                color = level_colors.get(val.strip(), DARK_TEXT)
                is_bold = val.strip() in level_colors
                add_cell_text(cell, val, font_size=9.5, bold=is_bold, color=color, align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                align = WD_ALIGN_PARAGRAPH.LEFT
                bold = (j == 0)
                col = INK if j == 0 else DARK_TEXT
                add_cell_text(cell, val, font_size=10, bold=bold, color=col, align=align)

    # 设置列宽（按页面可用宽度 15.7cm 分配）
    col_widths = [Cm(3.8), Cm(8.5), Cm(3.4)]
    for j, w in enumerate(col_widths):
        for row in table.rows:
            row.cells[j].width = w

    # 表格边框
    tbl = table._tbl
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'  <w:left w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'  <w:right w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'</w:tblBorders>'
    )
    tbl.tblPr.append(tblBorders)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_tier_table(headers, rows):
    """层级说明表（两列）"""
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        set_cell_shading(cell, GOLD_HEX)
        set_cell_margins(cell)
        add_cell_text(cell, h, font_size=11, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

    for i, row_data in enumerate(rows):
        bg = "FFFFFF" if i % 2 == 0 else "F8F6F1"
        for j, val in enumerate(row_data):
            cell = table.cell(i + 1, j)
            set_cell_shading(cell, bg)
            set_cell_margins(cell)
            bold = (j == 0)
            col = INK if j == 0 else DARK_TEXT
            add_cell_text(cell, val, font_size=10.5, bold=bold, color=col)

    col_widths = [Cm(3.5), Cm(12.2)]
    for j, w in enumerate(col_widths):
        for row in table.rows:
            row.cells[j].width = w

    tbl = table._tbl
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'  <w:left w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'  <w:right w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:color="{RULE_HEX}"/>'
        f'</w:tblBorders>'
    )
    tbl.tblPr.append(tblBorders)


# ============================================================
# 封面页
# ============================================================

# 空行留白
for _ in range(6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)

# 金色品牌条
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("━━━━━━━━━━━━")
run.font.size = Pt(12)
run.font.color.rgb = GOLD

# 产品名
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(4)
run = p.add_run("一码通")
run.font.size = Pt(42)
run.font.bold = True
run.font.color.rgb = INK
run.font.name = FONT_DISPLAY
rPr = run._element.find(qn('w:rPr'))
rFonts = rPr.find(qn('w:rFonts'))
rFonts.set(qn('w:eastAsia'), FONT_DISPLAY)

# 英文名
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
run = p.add_run("PRODUCT CAPABILITY GUIDE")
run.font.size = Pt(11)
run.font.color.rgb = GOLD
run.font.name = "Arial"

# 金色分隔线
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("━━━━━━━━━━━━")
run.font.size = Pt(12)
run.font.color.rgb = GOLD

# 副标题
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(24)
run = p.add_run("功能清单 · 客户沟通版")
run.font.size = Pt(18)
run.font.color.rgb = INK
run.font.name = FONT_DISPLAY
rPr = run._element.find(qn('w:rPr'))
rFonts = rPr.find(qn('w:rFonts'))
rFonts.set(qn('w:eastAsia'), FONT_DISPLAY)

# 定位语
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
run = p.add_run("包装扫码运营平台")
run.font.size = Pt(13)
run.font.color.rgb = MUTED
run.font.name = FONT_BODY
rPr = run._element.find(qn('w:rPr'))
rFonts = rPr.find(qn('w:rFonts'))
rFonts.set(qn('w:eastAsia'), FONT_BODY)

# 底部日期
for _ in range(8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("2026 年 8 月")
run.font.size = Pt(12)
run.font.color.rgb = MUTED
run.font.name = "Arial"

# 分页
doc.add_page_break()


# ============================================================
# 一、产品简介
# ============================================================

add_heading_1("一、产品简介")

add_body_paragraph(
    "一码通是一套面向食品、农产品及消费品品牌的包装扫码运营平台。企业可以把产品包装上的二维码升级为连接消费者的数字入口，在一次扫码中完成产品信息展示、批次溯源、轻量验真、权益领取、私域承接、复购跳转、数据分析和渠道风险预警。",
    indent_first=True
)

# 流程图（专业 PNG 图片）
add_heading_2("核心流程")

flow_img_path = "/tmp/flow_diagram_v2.png"
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(4)
run = p.add_run()
run.add_picture(flow_img_path, width=Cm(15.7))


# ============================================================
# 二、客户可以获得什么
# ============================================================

add_heading_1("二、客户可以获得什么")

customer_values = [
    ("信任增强 — ", "让包装二维码从静态展示工具变成长期可运营的消费者入口。"),
    ("产品背书 — ", "用产品、批次、检测报告和资质信息增强消费者信任。"),
    ("复购转化 — ", "通过权益、留资、企业微信和商城入口促进复购与私域沉淀。"),
    ("数据可视 — ", "看清不同产品、活动、码批次和渠道带来的扫码与转化结果。"),
    ("渠道预警 — ", "通过一物一码、扫码位置和渠道流向发现疑似复制码或窜货线索。"),
    ("灵活运营 — ", "由品牌方自主运营，或由服务商在明确授权范围内代配置、代运营。"),
]

for prefix, desc in customer_values:
    add_bullet(desc, bold_prefix=prefix)


# ============================================================
# 三、功能清单
# ============================================================

doc.add_page_break()
add_heading_1("三、功能清单")

headers = ["功能", "客户可实现的业务结果", "交付层级"]

# --- 3.1 品牌档案 ---
add_heading_2("1. 品牌、产品与批次数字档案")
add_feature_table(headers, [
    ["品牌资料管理", "统一维护品牌名称、Logo、介绍及对外展示信息", "基础能力"],
    ["产品管理", "管理产品名称、品类、产地、图片、详情和品牌故事", "基础能力"],
    ["SKU 管理", "按规格、包装类型、商品条码等区分具体商品", "基础能力"],
    ["生产批次管理", "维护批次号、生产日期、保质期、产地及关联 SKU", "基础能力"],
    ["产品资料附件", "展示检测报告、资质证书、图文、视频等信任资料", "基础能力"],
    ["产品关系管理", "建立品牌、产品、SKU、生产批次和码批次的关联", "基础能力"],
    ["批量资料导入", "批量导入产品、批次或其他主数据", "对接/定制"],
])

# --- 3.2 二维码 ---
add_heading_2("2. 二维码与一物一码管理")
add_feature_table(headers, [
    ["批次码", "同一批产品使用统一二维码，展示产品或批次信息", "基础能力"],
    ["一物一码", "每件商品使用唯一二维码，支持单件识别、首扫和重复扫码记录", "基础能力"],
    ["码批次管理", "按产品、生产批次或活动生成和管理一批二维码", "基础能力"],
    ["码包导出", "导出码文件，供标签打印、可变印刷或产线赋码使用", "基础能力"],
    ["码状态管理", "管理未激活、有效、冻结、作废等状态", "基础能力"],
    ["码查询与追踪", "查询单码所属产品、批次、状态及扫码记录", "基础能力"],
    ["已有码接管", "评估旧码中转、批量导入、域名接管或重新贴码等方案", "对接/定制"],
    ["外码/内码双码", "外码用于展示，内码用于开封后验真或权益领取", "增强能力"],
    ["印刷与产线赋码对接", "配合印刷厂、标签厂、喷码机或激光打码设备落码", "项目服务"],
])

# --- 3.3 H5 ---
add_heading_2("3. 消费者扫码 H5 页面")
add_feature_table(headers, [
    ["移动端扫码页面", "消费者使用微信、手机相机或浏览器扫码后直接打开", "基础能力"],
    ["页面模板", "快速套用产品展示、溯源、活动等页面结构", "基础能力"],
    ["模块化页面配置", "按需组合品牌、产品、批次、报告、权益、留资和跳转模块", "基础能力"],
    ["品牌视觉配置", "设置 Logo、品牌色、图片和品牌介绍", "基础能力"],
    ["页面预览", "发布前预览消费者实际看到的页面效果", "基础能力"],
    ["页面版本与发布", "保存页面版本，控制草稿和已发布内容", "基础能力"],
    ["自有域名", "使用客户自己的域名承载扫码页面", "对接/定制"],
    ["多语言展示", "面向不同市场配置多语言内容", "增强能力"],
])

# --- 3.4 溯源 ---
add_heading_2("4. 溯源与轻量验真")
add_feature_table(headers, [
    ["产品溯源展示", "消费者免登录查看产品、产地、生产批次和信任资料", "基础能力"],
    ["首次验证记录", "识别平台记录中的首次有效验证", "基础能力"],
    ["重复验证提示", "展示首次验证时间和累计验证次数，帮助消费者识别异常", "基础能力"],
    ["未激活码提示", "识别尚未启用或可能提前流出的码", "基础能力"],
    ["冻结码处理", "暂停敏感权益，同时保留必要的产品与溯源信息", "基础能力"],
    ["作废码处理", "停止正常验真和权益参与，并保留处置记录", "基础能力"],
    ["异常使用信号", "根据高频、异地或其他异常扫码形成风险提示", "增强能力"],
    ["风险人工处置", "由授权人员核查、记录原因并执行冻结、解冻等操作", "增强能力"],
])

add_note_block(
    "普通二维码可以证明平台中的码记录、产品资料和本次扫码信号，但不能单独证明实物绝对为正品。"
    "对于高价值商品，可结合内外双码、刮开层、易碎标签、NFC 等物理方案增强保护。"
)

# --- 3.5 营销活动 ---
add_heading_2("5. 营销活动与消费者权益")
add_feature_table(headers, [
    ["活动管理", "配置活动名称、时间、适用产品、页面和参与规则", "基础能力"],
    ["权益管理", "管理优惠券、礼品、积分或外部权益等营销内容", "基础能力"],
    ["权益库存", "设置总量、已使用量、剩余量和领取限制", "基础能力"],
    ["首扫活动", "针对符合条件的首次参与者发放指定权益", "基础能力"],
    ["领取记录", "查询消费者权益领取和发放状态", "基础能力"],
    ["重复领取控制", "按活动规则限制重复领取，避免重复发放", "基础能力"],
    ["权益履约与重试", "跟踪权益发放结果，对失败任务进行处理", "增强能力"],
    ["外部券码池/API 发券", "对接客户或第三方平台的优惠券和权益系统", "对接/定制"],
    ["活动风控", "对高频参与、异常设备或风险扫码暂停敏感权益", "增强能力"],
])

# --- 3.6 私域 ---
add_heading_2("6. 私域、留资与复购承接")
add_feature_table(headers, [
    ["企业微信入口", "引导消费者添加企业微信或进入客户群", "基础能力"],
    ["公众号/小程序入口", "跳转到客户已有公众号或小程序", "基础能力"],
    ["商城与电商店铺入口", "跳转客户现有商城、淘宝、京东、抖店等购买渠道", "基础能力"],
    ["留资表单", "在明确告知和授权后收集消费者联系方式或需求", "基础能力"],
    ["授权与同意记录", "保存授权场景、版本和时间，支持后续审计", "基础能力"],
    ["企业微信确认回传", "通过企业微信官方回调确认实际添加结果", "对接/定制"],
    ["订单数据回流", "通过标准文件或可信接口导入订单结果", "对接/定制"],
    ["CRM/会员系统同步", "将消费者线索或结果同步到客户现有系统", "对接/定制"],
])

add_note_block(
    '点击企业微信或商城入口属于\u201c转化意向\u201d；实际添加、领取成功或可信订单回流后，'
    '才计为确认转化。外部平台能力取决于客户账号权限和平台开放接口。'
)

# --- 3.7 会员 ---
add_heading_2("7. 会员与消费者运营")
add_feature_table(headers, [
    ["消费者档案", "汇总消费者的授权信息、扫码和参与记录", "基础能力"],
    ["匿名访客识别", "未登录时记录基础访问，不强制收集个人信息", "基础能力"],
    ["标签与分群", "按产品兴趣、活动参与和消费结果划分人群", "增强能力"],
    ["积分账户", "记录积分获得、使用和余额变化", "增强能力"],
    ["积分规则", "按扫码、活动或业务规则发放积分", "增强能力"],
    ["会员数据导出/同步", "在权限和合规控制下导出或同步消费者数据", "对接/定制"],
])

# --- 3.8 数据 ---
add_heading_2("8. 数据分析与经营复盘")
add_feature_table(headers, [
    ["经营总览", "查看扫码、访客、活动、权益和近期趋势", "基础能力"],
    ["扫码分析", "按时间、产品、码批次、地区等维度分析扫码表现", "基础能力"],
    ["活动分析", "查看访问、参与、领取及不同活动的转化表现", "基础能力"],
    ["意向与结果分层", "分开展示跳转点击与领取、添加、订单等确认结果", "基础能力"],
    ["订单与净成交额分析", "基于可信订单及退款数据统计订单和净成交额", "增强能力"],
    ["渠道归因", "分析不同渠道、码批次和入口带来的访问与结果", "增强能力"],
    ["数据质量提示", "区分未连接、同步为零、覆盖不完整和同步失败", "增强能力"],
    ["数据导出", "按权限导出业务数据并保留必要审计记录", "基础能力"],
])

# --- 3.9 渠道 ---
add_heading_2("9. 渠道管理与防窜线索")
add_feature_table(headers, [
    ["经销商管理", "建立经销商档案、账号和经营范围", "增强能力"],
    ["区域与门店管理", "管理授权区域、门店及渠道层级", "增强能力"],
    ["码批次流向登记", "记录某批货发往哪个经销商和授权区域", "增强能力"],
    ["扫码位置观察", "在消费者授权和合规前提下记录位置或区域信号", "增强能力"],
    ["跨区扫码观察", "对授权区域外的扫码形成待观察记录", "增强能力"],
    ["窜货线索聚合", "将持续偏离等信号聚合为可调查的业务线索", "增强能力"],
    ["经销商举证", "经销商可在自身权限范围内提交调货、订单或物流说明", "增强能力"],
    ["人工调查与结论", "分别记录调查进度及误报、正常调货、资料错误、确认窜货等结论", "增强能力"],
    ["线索审计", "保存处理人、原因、时间和历史记录", "增强能力"],
])

add_note_block(
    "系统根据扫码位置、渠道流向和信号质量形成风险观察或调查线索，"
    "不会仅凭一次低精度异地扫码自动认定窜货，也不会自动处罚或停用渠道。"
)

# --- 3.10 风控 ---
add_heading_2("10. 风控、权限与合规")
add_feature_table(headers, [
    ["多角色权限", "为管理员、运营、渠道、经销商、代运营等角色分配不同权限", "基础能力"],
    ["多组织账号", "按企业、部门或业务组织管理员工账号", "基础能力"],
    ["高风险操作控制", "对冻结、作废、敏感导出等操作保留原因和操作记录", "基础能力"],
    ["操作审计日志", "查询关键配置和业务操作的操作者及时间", "基础能力"],
    ["隐私政策配置", "配置消费者可见的隐私政策和授权说明", "基础能力"],
    ["个人信息保护", "对敏感数据采用授权、权限、加密及审计机制", "基础能力"],
    ["租户数据隔离", "不同客户的数据和账号相互隔离", "基础能力"],
    ["风险规则", "配置扫码、活动和渠道相关的风险识别规则", "增强能力"],
])

# --- 3.11 区域品牌 ---
add_heading_2("11. 区域品牌、协会与代运营")
add_feature_table(headers, [
    ["区域品牌组织", "建立区域品牌、协会或产业组织的管理结构", "增强能力"],
    ["成员企业管理", "下挂成员企业并管理成员状态", "增强能力"],
    ["统一模板与背书", "为成员企业提供统一品牌展示和页面模板", "对接/定制"],
    ["汇总数据看板", "汇总查看成员企业、产品和运营数据", "增强能力"],
    ["代运营工作台", "服务商集中管理客户任务和交付进度", "增强能力"],
    ["授权代运营", "由客户明确授权服务商可操作的租户和功能范围", "基础能力"],
    ["客户上线协助", "提供资料整理、页面配置、活动配置、码包生成和上线检查", "项目服务"],
    ["运营复盘服务", "按约定周期提供数据分析和优化建议", "项目服务"],
])

# --- 3.12 平台 ---
add_heading_2("12. 平台与外部集成")
add_feature_table(headers, [
    ["SaaS 客户开通", "为客户创建独立账号空间、套餐和使用额度", "基础能力"],
    ["套餐与额度管理", "管理产品数、码量、扫码量等服务额度", "基础能力"],
    ["客户健康与服务管理", "平台侧查看客户使用及服务状态", "增强能力"],
    ["短信/微信服务", "对接验证码、通知或微信相关能力", "对接/定制"],
    ["Webhook/Open API", "与客户现有系统交换业务数据", "对接/定制"],
    ["商城、CRM、ERP 对接", "根据客户系统和接口条件设计集成方案", "对接/定制"],
    ["私有化或专属部署", "根据客户安全、网络和运维要求单独评估", "对接/定制"],
])


# ============================================================
# 四、功能层级说明
# ============================================================

doc.add_page_break()
add_heading_1("四、功能层级说明")

add_tier_table(["层级", "说明"], [
    ["基础能力", "构成包装扫码、产品展示和运营闭环的通用能力，具体额度按所选套餐执行"],
    ["增强能力", "适合一物一码运营、会员、渠道、防窜或区域品牌等进阶场景，可按模块选配"],
    ["对接/定制", "依赖客户现有系统、外部平台接口、账号权限或个性化流程，需要评估后确定范围"],
    ["项目服务", "包括资料整理、印刷对接、页面配置、活动运营、培训和复盘等人工交付服务"],
])


# ============================================================
# 五、商务与交付说明
# ============================================================

add_heading_1("五、商务与交付说明")

commercial_notes = [
    "本清单用于售前能力沟通，不等同于最终合同功能范围。",
    "最终交付内容、服务额度、上线时间和费用，以双方确认的套餐、解决方案及合同为准。",
    "企业微信、电商、短信、CRM、ERP、印刷设备等外部能力，需要客户提供相应账号、接口权限或合作方支持。",
    "现金红包、支付、物流、售后和财务结算不属于标准 SaaS 闭环，如有需求需单独评估。",
    "涉及消费者个人信息的功能，应遵循合法、正当、必要原则，并结合客户实际业务完成隐私政策、授权和数据处理配置。",
]

for note in commercial_notes:
    add_bullet(note)


# ============================================================
# 页眉页脚
# ============================================================

for section in doc.sections:
    # 页眉
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)

    # 页眉文字
    run = hp.add_run("一码通 · 功能清单（客户沟通版）")
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    run.font.name = FONT_BODY
    rPr = run._element.find(qn('w:rPr'))
    rFonts = rPr.find(qn('w:rFonts'))
    rFonts.set(qn('w:eastAsia'), FONT_BODY)

    # 页眉底线（用段落下边框）
    pPr = hp._element.find(qn('w:pPr'))
    if pPr is None:
        pPr = parse_xml(f'<w:pPr {nsdecls("w")}></w:pPr>')
        hp._element.insert(0, pPr)
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="4" w:color="{GOLD_HEX}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)

    # 页脚
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)

    run = fp.add_run("2026 年 8 月")
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    run.font.name = "Arial"


# ============================================================
# 保存
# ============================================================

output_path = "/Users/ericding/code/agriculture/yimatong/docs/04_marketing_sales/一码通功能清单（客户沟通版）.docx"
doc.save(output_path)
print(f"✓ 文档已保存：{output_path}")
