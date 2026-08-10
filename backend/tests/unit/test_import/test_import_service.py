"""ERP 导入服务单元测试"""

import asyncio
import io
import uuid
from datetime import UTC, date, datetime, timedelta
from unittest.mock import AsyncMock

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.core.exceptions import QuotaExceededError
from app.models.plan import TenantQuotaUsage
from app.models.product import SKU, BatchStatus, Brand, Product, ProductionBatch
from app.models.tenant import Tenant
from app.services import import_service as import_service_module
from app.services.import_service import (
    SOURCE_SYSTEM,
    ExcelImportService,
    ImportResult,
    ParsedBatch,
    ParsedBrand,
    ParsedProduct,
    ParsedSKU,
)
from app.services.quota import QUOTA_RECONCILIATION_SOURCE_REVISION


@pytest.fixture
def tenant_id():
    return uuid.uuid4()


@pytest.fixture
def service():
    return ExcelImportService()


def _make_brand_sheet(wb, rows):
    ws = wb.create_sheet("品牌")
    ws.append(["品牌名称*", "品牌Logo URL", "品牌描述", "external_id"])
    for row in rows:
        ws.append(row)


def _make_product_sheet(wb, rows):
    ws = wb.create_sheet("产品")
    ws.append(["品牌(名称)*", "产品名称*", "产品分类", "产品描述", "external_id"])
    for row in rows:
        ws.append(row)


def _make_sku_sheet(wb, rows):
    ws = wb.create_sheet("SKU")
    ws.append(["产品(名称)*", "SKU编码*", "SKU名称*", "规格(JSON)", "external_id"])
    for row in rows:
        ws.append(row)


def _make_batch_sheet(wb, rows):
    ws = wb.create_sheet("批次")
    ws.append(["SKU(编码)*", "批次号*", "生产日期*", "有效期至*", "external_id", "产地"])
    for row in rows:
        ws.append(row)


class TestTemplateGeneration:
    def test_generate_template_has_all_sheets(self, service):
        content = service.generate_template()
        assert content is not None

        from openpyxl import load_workbook

        if isinstance(content, io.BytesIO):
            content = content.getvalue()
        wb = load_workbook(io.BytesIO(content))
        assert "品牌" in wb.sheetnames
        assert "产品" in wb.sheetnames
        assert "SKU" in wb.sheetnames
        assert "批次" in wb.sheetnames


class TestExcelParsing:
    async def test_xlsx_parse_capacity_rejects_same_tenant_without_thread_queue(self, monkeypatch):
        capacity = import_service_module.XLSXParseCapacity(global_limit=2, per_tenant_limit=1)
        service = ExcelImportService(parse_capacity=capacity)
        tenant_id = uuid.uuid4()
        lease = service.reserve_parse_capacity(tenant_id)
        to_thread = AsyncMock(return_value=ImportResult())
        monkeypatch.setattr(asyncio, "to_thread", to_thread)
        try:
            with pytest.raises(import_service_module.XLSXParseCapacityUnavailable):
                await service.parse_and_validate(b"workbook", tenant_id)
        finally:
            lease.release()

        to_thread.assert_not_awaited()

    async def test_xlsx_parse_capacity_preserves_slot_for_another_tenant(self, monkeypatch):
        capacity = import_service_module.XLSXParseCapacity(global_limit=2, per_tenant_limit=1)
        service = ExcelImportService(parse_capacity=capacity)
        first_tenant = uuid.uuid4()
        other_tenant = uuid.uuid4()
        lease = service.reserve_parse_capacity(first_tenant)
        to_thread = AsyncMock(return_value=ImportResult())
        monkeypatch.setattr(asyncio, "to_thread", to_thread)
        try:
            result = await service.parse_and_validate(b"workbook", other_tenant)
        finally:
            lease.release()

        assert result.errors == []
        to_thread.assert_awaited_once()

    async def test_xlsx_parse_capacity_is_globally_bounded(self):
        capacity = import_service_module.XLSXParseCapacity(global_limit=2, per_tenant_limit=1)
        service = ExcelImportService(parse_capacity=capacity)
        first = service.reserve_parse_capacity(uuid.uuid4())
        second = service.reserve_parse_capacity(uuid.uuid4())
        try:
            assert first is not None
            assert second is not None
            assert service.reserve_parse_capacity(uuid.uuid4()) is None
        finally:
            first.release()
            second.release()

    async def test_xlsx_parse_capacity_releases_after_parser_exception(self, monkeypatch):
        capacity = import_service_module.XLSXParseCapacity(global_limit=2, per_tenant_limit=1)
        service = ExcelImportService(parse_capacity=capacity)
        tenant_id = uuid.uuid4()
        to_thread = AsyncMock(side_effect=[RuntimeError("parser failed"), ImportResult()])
        monkeypatch.setattr(asyncio, "to_thread", to_thread)

        failed = await service.parse_and_validate(b"workbook", tenant_id)
        recovered = await service.parse_and_validate(b"workbook", tenant_id)

        assert [(error.sheet, error.row, error.message) for error in failed.errors] == [
            ("文件", 0, "无法解析 Excel 文件")
        ]
        assert recovered.errors == []
        assert to_thread.await_count == 2

    @pytest.mark.parametrize(
        ("limit_name", "limit_value", "message"),
        [
            ("MAX_XLSX_ENTRIES", 1, "Excel 文件包含过多压缩条目"),
            ("MAX_XLSX_TOTAL_UNCOMPRESSED", 1, "Excel 文件解压后总大小超过限制"),
            ("MAX_XLSX_SINGLE_UNCOMPRESSED", 1, "Excel 文件包含过大的压缩条目"),
            ("MAX_XLSX_COMPRESSION_RATIO", 1.0, "Excel 文件压缩比超过限制"),
        ],
    )
    async def test_xlsx_zip_preflight_limits(self, service, tenant_id, monkeypatch, limit_name, limit_value, message):
        content = service.generate_template().getvalue()
        monkeypatch.setattr(import_service_module, limit_name, limit_value, raising=False)

        result = await service.parse_and_validate(content, tenant_id)

        assert [(error.sheet, error.row, error.message) for error in result.errors] == [("文件", 0, message)]

    async def test_invalid_xlsx_error_is_sanitized(self, service, tenant_id):
        result = await service.parse_and_validate(b"not-an-xlsx", tenant_id)

        assert [(error.sheet, error.row, error.message) for error in result.errors] == [
            ("文件", 0, "无法解析 Excel 文件")
        ]

    async def test_openpyxl_parsing_runs_in_worker_thread(self, service, tenant_id, monkeypatch):
        content = service.generate_template().getvalue()
        calls = []

        async def recording_to_thread(func, *args, **kwargs):
            calls.append(func)
            return func(*args, **kwargs)

        monkeypatch.setattr(asyncio, "to_thread", recording_to_thread)

        result = await service.parse_and_validate(content, tenant_id)

        assert result.errors == []
        assert len(calls) == 1

    @pytest.mark.parametrize(
        ("limit_name", "limit_value", "message"),
        [
            ("MAX_XLSX_SHEETS", 1, "Excel Sheet 数量超过限制"),
            ("MAX_XLSX_ROWS_PER_SHEET", 1, "Excel Sheet 行数超过限制"),
            ("MAX_XLSX_CELLS", 1, "Excel 单元格数量超过限制"),
        ],
    )
    async def test_xlsx_workbook_shape_limits(self, service, tenant_id, monkeypatch, limit_name, limit_value, message):
        content = service.generate_template().getvalue()
        monkeypatch.setattr(import_service_module, limit_name, limit_value, raising=False)

        result = await service.parse_and_validate(content, tenant_id)

        assert [(error.sheet, error.row, error.message) for error in result.errors] == [("文件", 0, message)]

    async def test_xlsx_row_errors_are_capped(self, service, tenant_id, monkeypatch):
        wb = Workbook()
        ws = wb.active
        ws.title = "品牌"
        ws.append(["品牌名称*", "品牌Logo URL", "品牌描述", "external_id"])
        for _ in range(5):
            ws.append(["X" * 101, "", "", ""])
        buf = io.BytesIO()
        wb.save(buf)
        monkeypatch.setattr(import_service_module, "MAX_XLSX_ERRORS", 2, raising=False)

        result = await service.parse_and_validate(buf.getvalue(), tenant_id)

        assert len(result.errors) == 2

    async def test_parse_minimal_excel(self, service, db, tenant_id):
        wb = Workbook()
        wb.remove(wb.active)

        _make_brand_sheet(wb, [["测试品牌", "", "品牌描述", ""]])
        _make_product_sheet(wb, [["测试品牌", "测试产品", "零食", "产品描述", ""]])
        _make_sku_sheet(wb, [["测试产品", "SKU-001", "500g装", "{}", ""]])
        _make_batch_sheet(wb, [["SKU-001", "BH001", "2026-01-01", "2027-01-01", "", "黑龙江省五常市"]])

        buf = io.BytesIO()
        wb.save(buf)
        content = buf.getvalue()

        result = await service.parse_and_validate(content, tenant_id)
        assert len(result.errors) == 0
        assert len(result.brands) == 1
        assert len(result.products) == 1
        assert len(result.skus) == 1
        assert len(result.batches) == 1
        assert result.batches[0].origin == "黑龙江省五常市"

    async def test_parse_empty_required_field_skipped(self, service, db, tenant_id):
        wb = Workbook()
        wb.remove(wb.active)

        _make_brand_sheet(wb, [["", "", "", ""]])  # 品牌名称为空 → 跳过
        _make_product_sheet(wb, [])
        _make_sku_sheet(wb, [])
        _make_batch_sheet(wb, [])

        buf = io.BytesIO()
        wb.save(buf)
        content = buf.getvalue()

        result = await service.parse_and_validate(content, tenant_id)
        assert len(result.brands) == 0

    async def test_parse_with_external_id(self, service, db, tenant_id):
        wb = Workbook()
        wb.remove(wb.active)

        _make_brand_sheet(wb, [["测试品牌", "", "", "ext_brand_001"]])
        _make_product_sheet(wb, [])
        _make_sku_sheet(wb, [])
        _make_batch_sheet(wb, [])

        buf = io.BytesIO()
        wb.save(buf)
        content = buf.getvalue()

        result = await service.parse_and_validate(content, tenant_id)
        assert len(result.brands) == 1
        assert result.brands[0].external_id == "ext_brand_001"

    async def test_parse_multiple_brands(self, service, db, tenant_id):
        wb = Workbook()
        wb.remove(wb.active)

        _make_brand_sheet(
            wb,
            [
                ["品牌A", "", "描述A", "ext_a"],
                ["品牌B", "", "描述B", "ext_b"],
            ],
        )
        _make_product_sheet(wb, [])
        _make_sku_sheet(wb, [])
        _make_batch_sheet(wb, [])

        buf = io.BytesIO()
        wb.save(buf)
        content = buf.getvalue()

        result = await service.parse_and_validate(content, tenant_id)
        assert len(result.brands) == 2


class TestExcelExecutionQuota:
    async def test_created_product_is_reserved_once_and_upsert_update_is_not_recounted(self, service, db, tenant_id):
        db.add(
            Tenant(
                id=tenant_id,
                name="Excel 配额租户",
                slug=f"excel-quota-{tenant_id.hex[:8]}",
                quota={"max_products": 1},
            )
        )
        db.add(
            TenantQuotaUsage(
                tenant_id=tenant_id,
                reconciled_at=datetime.now(UTC),
                source_revision=QUOTA_RECONCILIATION_SOURCE_REVISION,
                enforcement_ready=True,
            )
        )
        parsed = ImportResult(
            brands=[ParsedBrand(2, "导入品牌", None, None, "brand-1")],
            products=[ParsedProduct(2, "导入品牌", "导入产品", None, None, "product-1")],
        )

        first = await service.execute_import(parsed, tenant_id, db)
        second = await service.execute_import(parsed, tenant_id, db)

        usage = await db.get(TenantQuotaUsage, tenant_id)
        assert usage is not None
        assert usage.products == 1
        assert first.products.created == 1
        assert second.products.updated == 1

    async def test_quota_error_aborts_excel_import_instead_of_becoming_row_error(self, service, db, tenant_id):
        db.add(
            Tenant(
                id=tenant_id,
                name="Excel 超额租户",
                slug=f"excel-overage-{tenant_id.hex[:8]}",
                quota={"max_products": 0},
            )
        )
        db.add(
            TenantQuotaUsage(
                tenant_id=tenant_id,
                reconciled_at=datetime.now(UTC),
                source_revision=QUOTA_RECONCILIATION_SOURCE_REVISION,
                enforcement_ready=True,
            )
        )
        parsed = ImportResult(
            brands=[ParsedBrand(2, "导入品牌", None, None, "brand-2")],
            products=[ParsedProduct(2, "导入品牌", "超额产品", None, None, "product-2")],
        )

        with pytest.raises(QuotaExceededError):
            await service.execute_import(parsed, tenant_id, db)


class TestExcelExecutionErrorSanitization:
    async def test_unexpected_error_returns_safe_code_reference_and_log(
        self, service, db, tenant_id, caplog, monkeypatch
    ):
        sensitive_exception = "SELECT secret FROM uploaded_value"
        monkeypatch.setattr(service, "_write_sync_record", AsyncMock(side_effect=RuntimeError(sensitive_exception)))
        caplog.set_level("WARNING", logger="app.services.import_service")

        report = await service.execute_import(
            ImportResult(
                brands=[
                    ParsedBrand(
                        row_num=2,
                        brand_name="敏感上传值",
                        logo_url=None,
                        description=None,
                        external_id=None,
                    )
                ]
            ),
            tenant_id,
            db,
        )

        error = report.brands.errors[0]
        assert report.brands.created == 0
        assert report.brands.updated == 0
        assert error.code == "IMPORT_ROW_FAILED"
        assert error.message == "该行导入失败，请核对数据后重试"
        assert error.reference_id.startswith("imp_")
        assert sensitive_exception not in error.message
        assert sensitive_exception not in caplog.text
        assert "敏感上传值" not in caplog.text
        assert error.reference_id in caplog.text

    async def test_integrity_error_returns_safe_code_reference_and_log(self, service, db, tenant_id, caplog):
        sensitive_value = "SENSITIVE-SKU-PARAMETER"
        brand = Brand(tenant_id=tenant_id, name="导入品牌")
        db.add(brand)
        await db.flush()
        product = Product(tenant_id=tenant_id, brand_id=brand.id, name="导入产品")
        db.add(product)
        await db.flush()
        db.add(SKU(tenant_id=tenant_id, product_id=product.id, code=sensitive_value, name="已有规格"))
        await db.flush()
        caplog.set_level("WARNING", logger="app.services.import_service")

        report = await service.execute_import(
            ImportResult(
                skus=[
                    ParsedSKU(
                        row_num=2,
                        product_name=product.name,
                        sku_code=sensitive_value,
                        sku_name="冲突规格",
                        specifications=None,
                        external_id=None,
                    )
                ]
            ),
            tenant_id,
            db,
        )

        assert len(report.skus.errors) == 1
        error = report.skus.errors[0]
        assert error.code == "IMPORT_ROW_CONFLICT"
        assert error.message == "该行数据冲突，未完成导入"
        assert error.reference_id.startswith("imp_")
        assert len(error.reference_id) == 36
        assert sensitive_value not in error.message
        assert sensitive_value not in caplog.text
        assert "INSERT INTO" not in caplog.text
        assert "parameters" not in caplog.text.lower()
        assert error.reference_id in caplog.text


class TestProductionBatchImportImmutability:
    async def test_excel_batch_origin_is_persisted(self, service, db, tenant_id):
        brand = Brand(tenant_id=tenant_id, name="导入品牌")
        db.add(brand)
        await db.flush()
        product = Product(tenant_id=tenant_id, brand_id=brand.id, name="导入产品")
        db.add(product)
        await db.flush()
        sku = SKU(tenant_id=tenant_id, product_id=product.id, code="ORIGIN-SKU", name="产地规格")
        db.add(sku)
        await db.flush()

        report = await service.execute_import(
            ImportResult(
                batches=[
                    ParsedBatch(
                        row_num=2,
                        sku_code=sku.code,
                        batch_code="ORIGIN-BATCH",
                        production_date=date.today(),
                        expiry_date=date.today() + timedelta(days=365),
                        external_id="origin-batch-1",
                        origin="黑龙江省哈尔滨市五常市",
                    )
                ]
            ),
            tenant_id,
            db,
        )

        batch = await db.scalar(
            select(ProductionBatch).where(
                ProductionBatch.tenant_id == tenant_id,
                ProductionBatch.external_id == "origin-batch-1",
            )
        )
        assert report.batches.total == 1
        assert report.batches.errors == []
        assert batch.origin == "黑龙江省哈尔滨市五常市"

    @pytest.mark.parametrize("status", [BatchStatus.recalled, BatchStatus.expired])
    async def test_non_active_batch_is_reported_and_not_updated(self, service, db, tenant_id, status):
        brand = Brand(tenant_id=tenant_id, name="导入品牌")
        db.add(brand)
        await db.flush()
        product = Product(tenant_id=tenant_id, brand_id=brand.id, name="导入产品")
        db.add(product)
        await db.flush()
        sku = SKU(tenant_id=tenant_id, product_id=product.id, code="IMMUTABLE-SKU", name="不可变规格")
        db.add(sku)
        await db.flush()
        batch = ProductionBatch(
            tenant_id=tenant_id,
            product_id=product.id,
            sku_id=sku.id,
            batch_code="ORIGINAL-BATCH",
            production_date=date(2026, 1, 1),
            expiry_date=date(2027, 1, 1),
            status=status,
            source_system=SOURCE_SYSTEM,
            external_id="immutable-batch-1",
            recall_reason="safety recall" if status == BatchStatus.recalled else None,
            recalled_at=datetime.now(UTC) if status == BatchStatus.recalled else None,
            recalled_by="test-actor" if status == BatchStatus.recalled else None,
        )
        db.add(batch)
        await db.flush()

        report = await service.execute_import(
            ImportResult(
                batches=[
                    ParsedBatch(
                        row_num=2,
                        sku_code=sku.code,
                        batch_code="MUST-NOT-CHANGE",
                        production_date=date(2026, 2, 1),
                        expiry_date=date(2027, 2, 1),
                        external_id="immutable-batch-1",
                    )
                ]
            ),
            tenant_id,
            db,
        )

        assert report.batches.created == 0
        assert report.batches.updated == 0
        assert [(error.code, error.message) for error in report.batches.errors] == [
            ("IMPORT_ROW_FAILED", "该行导入失败，请核对数据后重试")
        ]
        assert report.batches.errors[0].reference_id.startswith("imp_")
        assert batch.batch_code == "ORIGINAL-BATCH"
        assert batch.production_date == date(2026, 1, 1)

    async def test_active_batch_with_past_expiry_is_reported_and_not_updated(self, service, db, tenant_id):
        brand = Brand(tenant_id=tenant_id, name="导入品牌")
        db.add(brand)
        await db.flush()
        product = Product(tenant_id=tenant_id, brand_id=brand.id, name="导入产品")
        db.add(product)
        await db.flush()
        sku = SKU(tenant_id=tenant_id, product_id=product.id, code="PAST-EXPIRY-SKU", name="过期规格")
        db.add(sku)
        await db.flush()
        batch = ProductionBatch(
            tenant_id=tenant_id,
            product_id=product.id,
            sku_id=sku.id,
            batch_code="ORIGINAL-BATCH",
            production_date=date.today() - timedelta(days=30),
            expiry_date=date.today() - timedelta(days=1),
            status=BatchStatus.active,
            source_system=SOURCE_SYSTEM,
            external_id="past-expiry-batch-1",
        )
        db.add(batch)
        await db.flush()

        report = await service.execute_import(
            ImportResult(
                batches=[
                    ParsedBatch(
                        row_num=2,
                        sku_code=sku.code,
                        batch_code="MUST-NOT-CHANGE",
                        production_date=date.today(),
                        expiry_date=date.today() + timedelta(days=365),
                        external_id="past-expiry-batch-1",
                    )
                ]
            ),
            tenant_id,
            db,
        )

        assert report.batches.updated == 0
        assert [(error.code, error.message) for error in report.batches.errors] == [
            ("IMPORT_ROW_FAILED", "该行导入失败，请核对数据后重试")
        ]
        assert report.batches.errors[0].reference_id.startswith("imp_")
        assert batch.batch_code == "ORIGINAL-BATCH"
        assert batch.expiry_date == date.today() - timedelta(days=1)

    async def test_external_id_cannot_retarget_an_active_batch(self, service, db, tenant_id):
        brand = Brand(tenant_id=tenant_id, name="导入品牌")
        db.add(brand)
        await db.flush()
        original_product = Product(tenant_id=tenant_id, brand_id=brand.id, name="原产品")
        target_product = Product(tenant_id=tenant_id, brand_id=brand.id, name="目标产品")
        db.add_all([original_product, target_product])
        await db.flush()
        original_sku = SKU(tenant_id=tenant_id, product_id=original_product.id, code="ORIGINAL-SKU", name="原规格")
        target_sku = SKU(tenant_id=tenant_id, product_id=target_product.id, code="TARGET-SKU", name="目标规格")
        db.add_all([original_sku, target_sku])
        await db.flush()
        batch = ProductionBatch(
            tenant_id=tenant_id,
            product_id=original_product.id,
            sku_id=original_sku.id,
            batch_code="ORIGINAL-BATCH",
            production_date=date(2026, 1, 1),
            expiry_date=date(2027, 1, 1),
            source_system=SOURCE_SYSTEM,
            external_id="authoritative-batch-1",
        )
        db.add(batch)
        await db.flush()

        report = await service.execute_import(
            ImportResult(
                batches=[
                    ParsedBatch(
                        row_num=2,
                        sku_code=target_sku.code,
                        batch_code="MUST-NOT-RETARGET",
                        production_date=date(2026, 2, 1),
                        expiry_date=date(2027, 2, 1),
                        external_id="authoritative-batch-1",
                    )
                ]
            ),
            tenant_id,
            db,
        )

        assert report.batches.created == 0
        assert report.batches.updated == 0
        assert [(error.code, error.message) for error in report.batches.errors] == [
            ("IMPORT_ROW_FAILED", "该行导入失败，请核对数据后重试")
        ]
        assert report.batches.errors[0].reference_id.startswith("imp_")
        assert batch.product_id == original_product.id
        assert batch.sku_id == original_sku.id
        assert batch.batch_code == "ORIGINAL-BATCH"
